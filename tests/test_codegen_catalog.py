"""End-to-end tests for catalog-driven generated pipeline execution."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import polars as pl
from codegen import generate_polars_pipeline_script
from models import TargetSchema
from parser import discover_source_tables


def _run_generated_pipeline(
    tmp_path: Path,
    mapping_spec: dict,
    target_schema: TargetSchema,
) -> Path:
    pipeline_dir = tmp_path / "pipeline"
    output_dir = tmp_path / "output"
    source_dir = tmp_path / "sources"
    pipeline_dir.mkdir()
    output_dir.mkdir()
    source_dir.mkdir(exist_ok=True)
    script = generate_polars_pipeline_script(mapping_spec, target_schema)
    pipeline_path = pipeline_dir / "pipeline.py"
    pipeline_path.write_text(script.content, encoding="utf-8")
    (pipeline_dir / "mapping.json").write_text(
        json.dumps(mapping_spec),
        encoding="utf-8",
    )
    subprocess.run(
        [
            sys.executable,
            str(pipeline_path),
            "--source-folder",
            str(source_dir),
            "--output-folder",
            str(output_dir),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return output_dir


def test_generated_pipeline_honours_csv_catalog_settings(tmp_path: Path) -> None:
    """Execution should reuse detected encoding, delimiter, and header row."""
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    source_path = source_dir / "clients.csv"
    contents = (
        "Client export\r\n"
        "id;name;note\r\n"
        '1;André;"uses; delimiter"\r\n'
        "2;Zoë;\r\n"
    ).encode("cp1252")
    source_path.write_bytes(contents)
    catalog = discover_source_tables(
        contents,
        "csv",
        original_filename=source_path.name,
    )
    table = catalog.tables[0]
    schema = TargetSchema(
        client_code="test",
        tables=[
            {
                "name": "records",
                "columns": [
                    {"name": "record_id", "dtype": "Int64"},
                    {"name": "client_name", "dtype": "String"},
                ],
            }
        ],
    )
    mapping = {
        "target_schema_json": schema.model_dump(mode="json"),
        "source_catalogs": [catalog.model_dump(mode="json")],
        "columns": [
            {
                "target_table": "records",
                "target_column": "record_id",
                "source_columns": [
                    {
                        "source_table_id": table.source_table_id,
                        "source_column_id": table.columns[0].source_column_id,
                        "source_table": table.display_name,
                        "source_column": "id",
                    }
                ],
            },
            {
                "target_table": "records",
                "target_column": "client_name",
                "source_columns": [
                    {
                        "source_table_id": table.source_table_id,
                        "source_column_id": table.columns[1].source_column_id,
                        "source_table": table.display_name,
                        "source_column": "name",
                    }
                ],
            },
        ],
    }

    output_dir = _run_generated_pipeline(tmp_path, mapping, schema)

    result = pl.read_csv(output_dir / "records.csv")
    assert result.to_dicts() == [
        {"record_id": 1, "client_name": "André"},
        {"record_id": 2, "client_name": "Zoë"},
    ]


def test_generated_pipeline_disambiguates_identical_sheet_names(
    tmp_path: Path,
) -> None:
    """Stable table IDs should distinguish matching sheet names across workbooks."""
    from openpyxl import Workbook

    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    catalogs = []
    for filename, value in (("first.xlsx", "wrong"), ("second.xlsx", "chosen")):
        path = source_dir / filename
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["id", "value"])
        sheet.append([1, value])
        workbook.save(path)
        catalogs.append(
            discover_source_tables(
                path.read_bytes(),
                "xlsx",
                original_filename=filename,
            )
        )
    chosen = catalogs[1].tables[0]
    schema = TargetSchema(
        client_code="test",
        tables=[
            {
                "name": "records",
                "columns": [{"name": "selected_value", "dtype": "String"}],
            }
        ],
    )
    mapping = {
        "target_schema_json": schema.model_dump(mode="json"),
        "source_catalogs": [
            catalog.model_dump(mode="json") for catalog in catalogs
        ],
        "columns": [
            {
                "target_table": "records",
                "target_column": "selected_value",
                "source_columns": [
                    {
                        "source_table_id": chosen.source_table_id,
                        "source_column_id": chosen.columns[1].source_column_id,
                        "source_table": chosen.display_name,
                        "source_column": "value",
                    }
                ],
            }
        ],
    }

    output_dir = _run_generated_pipeline(tmp_path, mapping, schema)

    result = pl.read_csv(output_dir / "records.csv")
    assert result.to_dicts() == [{"selected_value": "chosen"}]

